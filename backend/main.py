from fastapi import FastAPI, Depends, HTTPException, status, UploadFile, File, WebSocket, WebSocketDisconnect
from fastapi.responses import StreamingResponse
from fastapi.security import OAuth2PasswordBearer, OAuth2PasswordRequestForm
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import or_ # <-- TAMBAHAN PENTING
from typing import Optional, List, Dict
import bcrypt
import jwt
import datetime
import json
import csv
import io
import re
import shutil
import uuid
import os
import random 
import string 
import openpyxl 
import google.generativeai as genai
from database import engine, get_db
import models
import schemas 

# --- KONFIGURASI AI ---
# ⚠️ PASTIKAN API KEY INI BENAR ⚠️
GOOGLE_API_KEY = "AIzaSyAMrTKGV3aMo2R0sAHzT-1Sd6sIts2CHm4" 

genai.configure(api_key=GOOGLE_API_KEY)
model = genai.GenerativeModel('gemini-2.0-flash')

SECRET_KEY = "rahasia_negara_on_cbt_jangan_disebar_bro"
ALGORITHM = "HS256"
ACCESS_TOKEN_EXPIRE_MINUTES = 60

models.Base.metadata.create_all(bind=engine)
oauth2_scheme = OAuth2PasswordBearer(tokenUrl="login")
oauth2_scheme_optional = OAuth2PasswordBearer(tokenUrl="login", auto_error=False)

app = FastAPI(title="ON-CBT Enterprise Final Fix 4", version="10.1.0")

origins = ["http://localhost:3000", "http://127.0.0.1:3000"]
app.add_middleware(CORSMiddleware, allow_origins=origins, allow_credentials=True, allow_methods=["*"], allow_headers=["*"])

os.makedirs("uploads", exist_ok=True)
app.mount("/static", StaticFiles(directory="uploads"), name="static")

# --- WEBSOCKET MANAGER ---
class ConnectionManager:
    def __init__(self): self.active_connections: Dict[int, Dict[str, WebSocket]] = {}
    async def connect(self, websocket: WebSocket, exam_id: int, username: str):
        await websocket.accept()
        if exam_id not in self.active_connections: self.active_connections[exam_id] = {}
        self.active_connections[exam_id][username] = websocket
    def disconnect(self, websocket: WebSocket, exam_id: int, username: str):
        if exam_id in self.active_connections:
            if username in self.active_connections[exam_id]: del self.active_connections[exam_id][username]
    async def broadcast(self, message: dict, exam_id: int):
        if exam_id in self.active_connections:
            for connection in self.active_connections[exam_id].values():
                try: await connection.send_json(message)
                except: pass
    async def send_personal(self, message: dict, exam_id: int, username: str):
        if exam_id in self.active_connections and username in self.active_connections[exam_id]:
            try: await self.active_connections[exam_id][username].send_json(message)
            except: pass
manager = ConnectionManager()

# --- HELPERS ---
def get_password_hash(p): return bcrypt.hashpw(p.encode(), bcrypt.gensalt()).decode()
def verify_password(p, h): return bcrypt.checkpw(p.encode(), h.encode())
def create_access_token(d): d.update({"exp": datetime.datetime.utcnow()+datetime.timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES)}); return jwt.encode(d, SECRET_KEY, ALGORITHM)
def get_current_user(token: str=Depends(oauth2_scheme), db: Session=Depends(get_db)):
    try: u = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM]).get("sub")
    except: raise HTTPException(401)
    user = db.query(models.User).options(joinedload(models.User.school)).filter(models.User.username == u).first()
    if not user: raise HTTPException(401)
    return user
def get_current_user_optional(token: str=Depends(oauth2_scheme_optional), db: Session=Depends(get_db)):
    if not token: return None
    try: u = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM]).get("sub"); return db.query(models.User).filter(models.User.username == u).first() if u else None
    except: return None
def check_is_teacher_or_admin(user: models.User):
    if user.role not in [models.UserRole.TEACHER, models.UserRole.ADMIN]: raise HTTPException(403, "Guru only")
def generate_random_password(length=6): return ''.join(random.choice(string.ascii_uppercase + string.digits) for _ in range(length))

# --- ENDPOINTS ---
@app.websocket("/ws/{exam_id}/{client_name}")
async def websocket_endpoint(websocket: WebSocket, exam_id: int, client_name: str):
    await manager.connect(websocket, exam_id, client_name)
    try:
        await manager.broadcast({"type":"status","student":client_name,"status":"online","message":"Bergabung"}, exam_id)
        while True:
            data = await websocket.receive_json()
            await manager.broadcast({"type":data.get("type"),"student":client_name,"status":data.get("status"),"message":data.get("message")}, exam_id)
    except WebSocketDisconnect:
        manager.disconnect(websocket, exam_id, client_name)
        await manager.broadcast({"type":"status","student":client_name,"status":"offline","message":"Terputus"}, exam_id)

@app.post("/schools/", response_model=schemas.SchoolResponse)
def create_school(s: schemas.SchoolCreate, db: Session=Depends(get_db)):
    if db.query(models.School).filter(models.School.code==s.code).first(): raise HTTPException(400, "Kode ada")
    ns = models.School(name=s.name, code=s.code, logo=s.logo); db.add(ns); db.commit(); db.refresh(ns); return ns

@app.post("/schools/logo")
async def upload_logo(file: UploadFile = File(...), db: Session = Depends(get_db), u: models.User = Depends(get_current_user)):
    check_is_teacher_or_admin(u)
    file_ext = file.filename.split(".")[-1]; filename = f"{uuid.uuid4()}.{file_ext}"
    with open(f"uploads/{filename}", "wb") as buffer: shutil.copyfileobj(file.file, buffer)
    logo_url = f"http://127.0.0.1:8000/static/{filename}"
    s = db.query(models.School).filter(models.School.id == u.school_id).first()
    if s: s.logo = logo_url; db.commit()
    return {"logo_url": logo_url}

@app.post("/classrooms/", response_model=schemas.ClassroomResponse)
def create_classroom(c: schemas.ClassroomCreate, db: Session=Depends(get_db), u: models.User=Depends(get_current_user)):
    check_is_teacher_or_admin(u)
    if db.query(models.Classroom).filter(models.Classroom.name==c.name, models.Classroom.school_id==u.school_id).first(): raise HTTPException(400, "Kelas ada")
    nc = models.Classroom(name=c.name, school_id=u.school_id); db.add(nc); db.commit(); db.refresh(nc); return nc

@app.get("/classrooms/", response_model=list[schemas.ClassroomResponse])
def get_classrooms(school_code: Optional[str]=None, db: Session=Depends(get_db), u: Optional[models.User]=Depends(get_current_user_optional)):
    if school_code:
        s = db.query(models.School).filter(models.School.code==school_code).first()
        return db.query(models.Classroom).filter(models.Classroom.school_id==s.id).all() if s else []
    if not u: raise HTTPException(401)
    return db.query(models.Classroom).filter(models.Classroom.school_id==u.school_id).all()

@app.get("/classrooms/{cid}/students", response_model=list[schemas.UserResponse])
def get_students(cid: int, db: Session=Depends(get_db), u: models.User=Depends(get_current_user)):
    check_is_teacher_or_admin(u)
    return db.query(models.User).filter(models.User.classroom_id==cid).all()

@app.post("/classrooms/{classroom_id}/import-students")
async def import_students_to_class(classroom_id: int, file: UploadFile = File(...), db: Session = Depends(get_db), u: models.User = Depends(get_current_user)):
    check_is_teacher_or_admin(u)
    if not file.filename.endswith('.xlsx'): raise HTTPException(400, "Format .xlsx")
    try:
        contents = await file.read(); wb = openpyxl.load_workbook(io.BytesIO(contents)); sheet = wb.active
        count = 0
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row[0]: continue
            full_name = str(row[0]); nisn = str(row[1]) if row[1] else str(random.randint(100000, 999999))
            username = nisn; raw_password = "123456" 
            if db.query(models.User).filter(models.User.username == username).first(): continue
            ns = models.User(username=username, full_name=full_name, role=models.UserRole.STUDENT, hashed_password=get_password_hash(raw_password), original_password=raw_password, school_id=u.school_id, classroom_id=classroom_id, nisn=nisn, exam_number=f"U-{nisn}")
            db.add(ns); count += 1
        db.commit(); return {"message": f"Sukses generate {count} akun siswa!"}
    except Exception as e: raise HTTPException(500, f"Error: {str(e)}")

@app.post("/users/", response_model=schemas.UserResponse)
def create_user(u: schemas.UserCreate, db: Session=Depends(get_db)):
    s = db.query(models.School).filter(models.School.code==u.school_code).first()
    if not s: raise HTTPException(404, "Sekolah 404")
    if db.query(models.User).filter(models.User.username==u.username).first(): raise HTTPException(400, "Username ada")
    nu = models.User(username=u.username, full_name=u.full_name, role=u.role, hashed_password=get_password_hash(u.password), school_id=s.id, classroom_id=u.classroom_id)
    db.add(nu); db.commit(); db.refresh(nu); return nu

@app.post("/login", response_model=schemas.Token)
def login(f: OAuth2PasswordRequestForm=Depends(), db: Session=Depends(get_db)):
    u = db.query(models.User).filter(models.User.username==f.username).first()
    if not u or not verify_password(f.password, u.hashed_password): raise HTTPException(401)
    return {"access_token": create_access_token({"sub": u.username}), "token_type": "bearer"}

@app.get("/users/me", response_model=schemas.UserResponse)
def me(u: models.User=Depends(get_current_user)): return u

# --- UJIAN ---
@app.post("/exams/", response_model=schemas.ExamResponse)
def create_exam(e: schemas.ExamCreate, db: Session=Depends(get_db), u: models.User=Depends(get_current_user)):
    check_is_teacher_or_admin(u)
    ne = models.Exam(title=e.title, description=e.description, duration_minutes=e.duration_minutes, created_by=u.id, school_id=u.school_id, is_active=True)
    if e.classroom_ids: ne.classrooms = db.query(models.Classroom).filter(models.Classroom.id.in_(e.classroom_ids)).all()
    db.add(ne); db.commit(); db.refresh(ne); return ne

# [🔥 FIX LOGIKA FILTER UJIAN 🔥]
@app.get("/exams/", response_model=list[schemas.ExamResponse])
def read_exams(skip: int=0, limit: int=100, db: Session=Depends(get_db), u: models.User=Depends(get_current_user)):
    # Ambil semua ujian di sekolah ini, bawa data soal dan kelas
    q = db.query(models.Exam).options(
        joinedload(models.Exam.classrooms), 
        joinedload(models.Exam.questions)
    ).filter(models.Exam.school_id==u.school_id)
    
    # Kalau SISWA & punya kelas:
    # Tampilkan jika (Ujian ini untuk kelas dia) ATAU (Ujian ini tidak punya target kelas/Global)
    if u.role == models.UserRole.STUDENT and u.classroom_id:
        q = q.filter(
            or_(
                models.Exam.classrooms.any(models.Classroom.id == u.classroom_id),
                ~models.Exam.classrooms.any() # <-- Ini kuncinya: Ujian 'Global' tetap muncul
            )
        )
    return q.offset(skip).limit(limit).all()

@app.delete("/exams/{exam_id}")
def delete_exam(exam_id: int, db: Session=Depends(get_db), u: models.User=Depends(get_current_user)):
    check_is_teacher_or_admin(u)
    ex = db.query(models.Exam).filter(models.Exam.id==exam_id).first()
    if not ex or ex.school_id!=u.school_id: raise HTTPException(404)
    db.query(models.Question).filter(models.Question.exam_id==exam_id).delete()
    db.query(models.ExamResult).filter(models.ExamResult.exam_id==exam_id).delete()
    db.delete(ex); db.commit(); return {"msg": "Deleted"}

@app.post("/exams/{exam_id}/questions/", response_model=schemas.QuestionResponse)
def create_q(exam_id: int, q: schemas.QuestionCreate, db: Session=Depends(get_db), u: models.User=Depends(get_current_user)):
    check_is_teacher_or_admin(u)
    nq = models.Question(**q.model_dump(), exam_id=exam_id); db.add(nq); db.commit(); db.refresh(nq); return nq

@app.post("/exams/{exam_id}/generate-ai")
def gen_ai(exam_id: int, r: schemas.GenerateRequest, db: Session=Depends(get_db), u: models.User=Depends(get_current_user)):
    check_is_teacher_or_admin(u)
    try:
        p = f"Buat {r.count} soal PG topik '{r.topic}'. JSON Array only: [{{'content':'...','options':{{'A':'...','B':'...'}},'correct_answer':'A','score':10}}]"
        res = model.generate_content(p)
        json_match = re.search(r"\[.*\]", res.text, re.DOTALL)
        if not json_match: raise ValueError("AI Format Error")
        data = json.loads(json_match.group(0))
        saved = []
        for i in data:
            nq = models.Question(exam_id=exam_id, content=i["content"], question_type="multiple_choice", options=json.dumps(i["options"]), correct_answer=i["correct_answer"], score=i.get("score",10))
            db.add(nq); saved.append(nq)
        db.commit(); return {"data": s}
    except Exception as e: raise HTTPException(500, str(e))

@app.get("/downloads/template")
def download_template():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Template Soal"
    ws.append(["Pertanyaan", "Opsi A", "Opsi B", "Opsi C", "Opsi D", "Kunci Jawaban (A/B/C/D)", "Skor"])
    ws.append(["Siapa presiden pertama RI?", "Soeharto", "Habibie", "Soekarno", "Jokowi", "C", 10])
    buffer = io.BytesIO(); wb.save(buffer); buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=template_soal_oncbt.xlsx"})

@app.post("/exams/{exam_id}/import")
async def import_questions(exam_id: int, file: UploadFile = File(...), db: Session = Depends(get_db), u: models.User = Depends(get_current_user)):
    check_is_teacher_or_admin(u)
    if not file.filename.endswith('.xlsx'): raise HTTPException(400, "Format harus .xlsx")
    try:
        contents = await file.read(); wb = openpyxl.load_workbook(io.BytesIO(contents)); sheet = wb.active
        saved_questions = []
        for idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True)):
            if not row or all(cell is None for cell in row): continue
            safe_row = list(row) + [None] * (7 - len(row))
            content = str(safe_row[0]) if safe_row[0] else None
            if not content: continue
            options = json.dumps({"A": str(safe_row[1] or ""), "B": str(safe_row[2] or ""), "C": str(safe_row[3] or ""), "D": str(safe_row[4] or "")})
            correct = str(safe_row[5]).strip().upper() if safe_row[5] else "A"
            if correct not in ["A", "B", "C", "D"]: correct = "A"
            try: score = int(safe_row[6]) if safe_row[6] is not None else 10
            except: score = 10
            nq = models.Question(exam_id=exam_id, content=content, question_type="multiple_choice", options=options, correct_answer=correct, score=score)
            db.add(nq); saved_questions.append(nq)
        db.commit()
        return {"message": f"Berhasil mengimport {len(saved_questions)} soal! Silakan cek daftar."}
    except Exception as e: print(f"Error import: {e}"); raise HTTPException(500, f"Gagal proses file: {str(e)}")

@app.post("/exams/{exam_id}/submit", response_model=schemas.ExamResultResponse)
def submit(exam_id: int, sub: schemas.AnswerSubmit, db: Session=Depends(get_db), u: models.User=Depends(get_current_user)):
    qs = db.query(models.Question).filter(models.Question.exam_id == exam_id).all()
    score = 0; saved_detail = {}
    for q in qs:
        ans = sub.answers.get(str(q.id))
        if ans is None: ans = sub.answers.get(q.id)
        if ans is None: ans = ""
        pts = 0
        if q.question_type == "multiple_choice" and isinstance(ans, str) and ans.lower() == q.correct_answer.lower(): pts = q.score
        elif q.question_type == "essay" and ans:
            try:
                r = model.generate_content(f"Soal:{q.content}\nKunci:{q.correct_answer}\nJwb:{ans}\nSkor 0.0-1.0:")
                m = re.search(r"(\d+(\.\d+)?)", r.text)
                ratio = float(m.group(1)) if m else 0.0
                pts = int(min(max(ratio,0.0),1.0) * q.score)
            except: pts = 0
        score += pts
        saved_detail[str(q.id)] = {"answer": ans, "score": pts, "max_score": q.score}
    res = models.ExamResult(exam_id=exam_id, student_id=u.id, score=score, answers=json.dumps(saved_detail))
    db.add(res); db.commit(); db.refresh(res); return res

@app.get("/exams/{exam_id}/results", response_model=list[schemas.ExamResultWithStudent])
def get_res(exam_id: int, db: Session=Depends(get_db), u: models.User=Depends(get_current_user)):
    check_is_teacher_or_admin(u)
    return db.query(models.ExamResult).options(joinedload(models.ExamResult.student)).filter(models.ExamResult.exam_id == exam_id).all()

class ExamResultDetailFull(schemas.ExamResultDetail):
    answers: Optional[str] = None 

@app.get("/results/{result_id}", response_model=ExamResultDetailFull) 
def get_result_detail(result_id: int, db: Session = Depends(get_db), u: models.User = Depends(get_current_user)):
    check_is_teacher_or_admin(u)
    res = db.query(models.ExamResult).options(joinedload(models.ExamResult.student),joinedload(models.ExamResult.exam).joinedload(models.Exam.questions)).filter(models.ExamResult.id == result_id).first()
    if not res: raise HTTPException(404, "Data tidak ditemukan")
    return res

@app.put("/results/{result_id}/score")
def update_score(result_id: int, req: schemas.ScoreUpdate, db: Session = Depends(get_db), u: models.User = Depends(get_current_user)):
    check_is_teacher_or_admin(u)
    res = db.query(models.ExamResult).filter(models.ExamResult.id == result_id).first()
    if not res: raise HTTPException(404, "Data tidak ditemukan")
    res.score = req.score; db.commit()
    return {"message": "Updated"}

@app.get("/exams/{exam_id}/export")
def exp_res(exam_id: int, db: Session=Depends(get_db), u: models.User=Depends(get_current_user)):
    check_is_teacher_or_admin(u)
    rs = db.query(models.ExamResult).options(joinedload(models.ExamResult.student)).filter(models.ExamResult.exam_id == exam_id).all()
    o = io.StringIO(); w = csv.writer(o)
    w.writerow(["No", "Nama", "Nilai"]); [w.writerow([i+1, r.student.full_name, r.score]) for i, r in enumerate(rs)]
    o.seek(0); return StreamingResponse(io.BytesIO(o.getvalue().encode()), media_type="text/csv", headers={"Content-Disposition": "attachment; filename=nilai.csv"})

@app.get("/exams/{exam_id}/analysis")
def analyze(exam_id: int, db: Session=Depends(get_db), u: models.User=Depends(get_current_user)):
    check_is_teacher_or_admin(u)
    qs = db.query(models.Question).filter(models.Question.exam_id == exam_id).all()
    rs = db.query(models.ExamResult).filter(models.ExamResult.exam_id == exam_id).all()
    if not rs: return {"data": []}
    data = []
    for q in qs:
        earned = 0
        for r in rs:
            det = json.loads(r.answers).get(str(q.id))
            if isinstance(det, dict): earned += det.get("score", 0)
            elif isinstance(det, str) and det.lower() == q.correct_answer.lower(): earned += q.score
        diff = earned / (q.score * len(rs)) if q.score > 0 else 0
        lbl, clr = ("Mudah", "green") if diff >= 0.7 else ("Sukar", "red") if diff <= 0.3 else ("Sedang", "yellow")
        data.append({"question_id": q.id, "content": q.content[:50], "correct_count": round(earned/q.score, 1), "difficulty_index": round(diff, 2), "label": lbl, "color": clr})
    return {"data": data}

@app.get("/exams/{exam_id}/attendance")
def get_exam_attendance(exam_id: int, db: Session = Depends(get_db), u: models.User = Depends(get_current_user)):
    check_is_teacher_or_admin(u)
    exam = db.query(models.Exam).options(joinedload(models.Exam.classrooms), joinedload(models.Exam.school)).filter(models.Exam.id == exam_id).first()
    if not exam: raise HTTPException(404, "Ujian tidak ditemukan")
    target_class_ids = [c.id for c in exam.classrooms]
    if not target_class_ids:
        students = db.query(models.User).filter(models.User.school_id == u.school_id, models.User.role == models.UserRole.STUDENT).all()
    else:
        students = db.query(models.User).options(joinedload(models.User.classroom)).filter(models.User.classroom_id.in_(target_class_ids), models.User.role == models.UserRole.STUDENT).all()
    results = db.query(models.ExamResult).filter(models.ExamResult.exam_id == exam_id).all()
    submitted_student_ids = [r.student_id for r in results]
    attendance_list = []
    for s in students:
        status = "Hadir" if s.id in submitted_student_ids else "Tidak Hadir"
        attendance_list.append({"student_id": s.id, "username": s.username, "full_name": s.full_name, "class_name": s.classroom.name if s.classroom else "-", "status": status})
    return {"exam_title": exam.title, "exam_date": exam.start_time or datetime.datetime.now(), "school_name": exam.school.name if exam.school else "Sekolah", "school_logo": exam.school.logo, "students": attendance_list}

@app.get("/student/my-results", response_model=list[schemas.ExamResultWithExam])
def my_res(db: Session=Depends(get_db), u: models.User=Depends(get_current_user)):
    return db.query(models.ExamResult).options(joinedload(models.ExamResult.exam)).filter(models.ExamResult.student_id == u.id).order_by(models.ExamResult.submitted_at.desc()).all()

# --- REMOTE CONTROL ENDPOINTS ---
@app.post("/exams/{exam_id}/reset/{student_id}")
async def reset_student_exam(exam_id: int, student_id: int, db: Session = Depends(get_db), u: models.User = Depends(get_current_user)):
    check_is_teacher_or_admin(u)
    db.query(models.ExamResult).filter(models.ExamResult.exam_id == exam_id, models.ExamResult.student_id == student_id).delete()
    db.commit()
    student = db.query(models.User).filter(models.User.id == student_id).first()
    if student:
        await manager.send_personal({"type": "command", "action": "reset_exam", "message": "Ujian Anda di-reset oleh Guru."}, exam_id, student.username)
    return {"message": "Siswa berhasil di-reset"}

@app.post("/exams/{exam_id}/add-time/{student_id}")
async def add_student_time(exam_id: int, student_id: int, minutes: int = 10, db: Session = Depends(get_db), u: models.User = Depends(get_current_user)):
    check_is_teacher_or_admin(u)
    student = db.query(models.User).filter(models.User.id == student_id).first()
    if student:
        await manager.send_personal({"type": "command", "action": "add_time", "minutes": minutes, "message": f"Waktu ditambah {minutes} menit."}, exam_id, student.username)
    return {"message": f"Waktu siswa ditambah {minutes} menit"}